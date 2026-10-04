# -*- coding: utf-8 -*-
"""Deck — template handling (layouts, slide factory, notes, metadata) for the DGA Arabic template."""
import copy
import datetime
import io
import os
import re
import zipfile
from xml.sax.saxutils import escape

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

from deck_lib import (E, NS, NSDECL, NS_A, NS_P, NS_R, F_SEMI, F_MED, F_LIGHT, NAVY, PURPLE, TEAL,
                      GREEN, WHITE, BODY, MUTED, Icons, Sl, para, run, sp_xml)

SKILL_DIR = os.environ.get(
    "DGA_SKILL_DIR",
    "/root/.claude/skills/synced/b505a9e7-16b3-408b-bb9a-326ffa85f894_2d75fd7b-5f81-49d3-965c-2b4232bf6942/dga-pptx-designer",
)
TEMPLATE = os.path.join(SKILL_DIR, "assets", "templates", "DGA Template Ar.pptx")

# geometry (inches)
SW, SH = 13.333, 7.5
R_EDGE, L_EDGE = 12.42, 0.92
CW = R_EDGE - L_EDGE
Y_BODY, Y_BODY_END = 2.20, 6.55

CLASSIFICATION = "عام - Public"
DECK_SHORT_TITLE = "التدريب في الذكاء الاصطناعي  |  تحديث لمعالي المحافظ"


def qn(tag):
    pfx, name = tag.split(":")
    return "{%s}%s" % (NS[pfx], name)


class Deck:
    def __init__(self, icons_dir, template=TEMPLATE):
        self.prs = Presentation(template)
        self.icons = Icons(icons_dir)
        self.slides = []          # list of dicts: {sl, kind, title, label}
        self._strip_template_slides()
        self._prep_layouts()

    # ------------------------------------------------------------------ template prep
    def _layout(self, partname):
        for m in self.prs.slide_masters:
            for l in m.slide_layouts:
                if str(l.part.partname) == partname:
                    return l
        raise KeyError(partname)

    def _strip_template_slides(self):
        lst = self.prs.slides._sldIdLst
        for sld in list(lst):
            self.prs.part.drop_rel(sld.rId)
            lst.remove(sld)
        # stale revision tracking from the template
        for rid, rel in list(self.prs.part.rels.items()):
            if rel.reltype.endswith("/revisionInfo") or "changesInfo" in rel.reltype:
                self.prs.part.drop_rel(rid)

    def _prep_layouts(self):
        # ---- content layout (Title Slide / slideLayout1): fixed title + lede + footer
        self.L_CONTENT = self._layout("/ppt/slideLayouts/slideLayout1.xml")
        tree = self.L_CONTENT.shapes._spTree
        title_sp = None
        for sp in tree.findall(qn("p:sp")):
            ph = sp.find(".//p:ph", NS)
            if ph is not None and ph.get("type") == "title":
                title_sp = sp
        spPr = title_sp.find(qn("p:spPr"))
        for ch in list(spPr):
            spPr.remove(ch)
        spPr.append(etree.fromstring(
            f'<a:xfrm xmlns:a="{NS_A}"><a:off x="{E(0.92)}" y="{E(0.40)}"/><a:ext cx="{E(11.5)}" cy="{E(1.10)}"/></a:xfrm>'))
        spPr.append(etree.fromstring(f'<a:prstGeom xmlns:a="{NS_A}" prst="rect"><a:avLst/></a:prstGeom>'))
        old = title_sp.find(qn("p:txBody"))
        title_sp.replace(old, etree.fromstring(
            f'<p:txBody {NSDECL}><a:bodyPr lIns="0" tIns="0" rIns="0" bIns="0" rtlCol="0" anchor="b"><a:noAutofit/></a:bodyPr>'
            f'<a:lstStyle><a:lvl1pPr marL="0" indent="0" algn="r" rtl="1"><a:lnSpc><a:spcPts val="3700"/></a:lnSpc>'
            f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:buNone/>'
            f'<a:defRPr sz="2800" b="1" i="0" kern="1200"><a:solidFill><a:srgbClr val="{PURPLE}"/></a:solidFill>'
            f'<a:latin typeface="{F_SEMI}" pitchFamily="2" charset="-78"/><a:ea typeface="+mj-ea"/>'
            f'<a:cs typeface="{F_SEMI}" pitchFamily="2" charset="-78"/></a:defRPr></a:lvl1pPr></a:lstStyle>'
            f'<a:p><a:r><a:rPr lang="ar-SA"/><a:t>عنوان الشريحة: رسالة واحدة واضحة</a:t></a:r></a:p></p:txBody>'))
        nxt = max(int(e.get("id")) for e in tree.iter() if e.tag == qn("p:cNvPr")) + 1
        lede = etree.fromstring(
            f'<p:sp {NSDECL}><p:nvSpPr><p:cNvPr id="{nxt}" name="Lede Placeholder"/>'
            f'<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
            f'<p:nvPr><p:ph type="body" sz="quarter" idx="13" hasCustomPrompt="1"/></p:nvPr></p:nvSpPr>'
            f'<p:spPr><a:xfrm><a:off x="{E(0.92)}" y="{E(1.58)}"/><a:ext cx="{E(11.5)}" cy="{E(0.46)}"/></a:xfrm>'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
            f'<p:txBody><a:bodyPr lIns="0" tIns="0" rIns="0" bIns="0" rtlCol="0" anchor="t"><a:noAutofit/></a:bodyPr>'
            f'<a:lstStyle><a:lvl1pPr marL="0" indent="0" algn="r" rtl="1"><a:lnSpc><a:spcPts val="2400"/></a:lnSpc>'
            f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:buNone/>'
            f'<a:defRPr sz="1600" b="0" i="0" kern="1200"><a:solidFill><a:srgbClr val="{TEAL}"/></a:solidFill>'
            f'<a:latin typeface="{F_MED}" pitchFamily="2" charset="-78"/><a:ea typeface="+mn-ea"/>'
            f'<a:cs typeface="{F_MED}" pitchFamily="2" charset="-78"/></a:defRPr></a:lvl1pPr></a:lstStyle>'
            f'<a:p><a:pPr lvl="0"/><a:r><a:rPr lang="ar-SA"/><a:t>الجملة الداعمة للرسالة</a:t></a:r></a:p></p:txBody></p:sp>')
        tree.append(lede)
        # footer texts live on the layout (edit once in Slide Master)
        for k, (name, x, w, algn, sz, color, font, txt) in enumerate([
            ("Classification", 5.17, 3.0, "ctr", 12, GREEN, F_SEMI, CLASSIFICATION),
            ("Deck short title", 7.9, 4.4, "r", 10, MUTED, F_MED, DECK_SHORT_TITLE),
        ]):
            nid = nxt + 1 + k
            tree.append(etree.fromstring(sp_xml(
                nid, name, x, 7.0, w, 0.36,
                [para(txt, sz=sz, color=color, font=font, algn=algn, pitch=sz * 1.3)],
                anchor="ctr", ins=(0, 0, 0, 0), txbox=True)))
        # ---- cover / divider layouts
        self.L_COVER = self._layout("/ppt/slideLayouts/slideLayout44.xml")     # network globe
        self.L_DIVIDER = self._layout("/ppt/slideLayouts/slideLayout43.xml")   # dark gradient + chart ghost

    # ------------------------------------------------------------------ slide factory
    def _fill_ph(self, slide, ph_type=None, idx=None, text="", rpr_extra=""):
        for sp in slide.shapes._spTree.findall(qn("p:sp")):
            ph = sp.find(".//p:ph", NS)
            if ph is None:
                continue
            if (ph_type and ph.get("type") == ph_type) or (idx is not None and ph.get("idx") == str(idx)):
                tx = sp.find(qn("p:txBody"))
                for ch in list(tx):
                    tx.remove(ch)
                tx.append(etree.fromstring(f'<a:bodyPr xmlns:a="{NS_A}"/>'))
                tx.append(etree.fromstring(f'<a:lstStyle xmlns:a="{NS_A}"/>'))
                segs = text.split("\n")
                inner = ""
                for k, seg in enumerate(segs):
                    if k:
                        inner += '<a:br><a:rPr lang="ar-SA" dirty="0"/></a:br>'
                    inner += (f'<a:r><a:rPr lang="ar-SA" dirty="0" noProof="1"{rpr_extra}/>'
                              f'<a:t>{escape(seg)}</a:t></a:r>')
                tx.append(etree.fromstring(f'<a:p xmlns:a="{NS_A}">{inner}</a:p>'))
                return sp
        return None

    def _sldnum(self, sl, idx=11):
        nid = sl.nid()
        sl._append(
            f'<p:sp {NSDECL}><p:nvSpPr><p:cNvPr id="{nid}" name="Slide Number"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
            f'<p:nvPr><p:ph type="sldNum" sz="quarter" idx="{idx}"/></p:nvPr></p:nvSpPr><p:spPr/>'
            f'<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:fld id="{{E783386A-6F2F-1D4B-9C0A-84490692B308}}" type="slidenum">'
            f'<a:rPr lang="en-US"/><a:t>‹#›</a:t></a:fld><a:endParaRPr lang="en-US" dirty="0"/></a:p></p:txBody></p:sp>')

    def content(self, title, lede=None, label="", kind="content"):
        slide = self.prs.slides.add_slide(self.L_CONTENT)
        sl = Sl(slide, self.icons, label=label or title[:24])
        self._fill_ph(slide, ph_type="title", text=title)
        if lede:
            self._fill_ph(slide, idx=13, text=lede)
        else:  # remove empty lede placeholder
            for sp in slide.shapes._spTree.findall(qn("p:sp")):
                ph = sp.find(".//p:ph", NS)
                if ph is not None and ph.get("idx") == "13":
                    slide.shapes._spTree.remove(sp)
        self._sldnum(sl)
        self.slides.append(dict(sl=sl, kind=kind, title=title, label=sl.label))
        return sl

    def source(self, sl, text, x=1.75):
        """small source / status line above the footer (left of the DGA corner mark)"""
        pr = para(text, sz=10.5, color=MUTED, algn="r", pitch=14)
        if pr.lines(R_EDGE - x) > 1:
            sl.warn(f"source line wraps to {pr.lines(R_EDGE - x)} lines")
        sl.text("المرجع", x, 6.60, R_EDGE - x, 0.38, [pr], anchor="t")

    def cover(self, program, title, subtitle, byline, label="cover"):
        slide = self.prs.slides.add_slide(self.L_COVER)
        sl = Sl(slide, self.icons, label=label)
        # layout title placeholder -> our big white title
        sp = self._fill_ph(slide, ph_type="title", text=title)
        spPr = sp.find(qn("p:spPr"))
        for ch in list(spPr):
            spPr.remove(ch)
        spPr.append(etree.fromstring(
            f'<a:xfrm xmlns:a="{NS_A}"><a:off x="{E(4.9)}" y="{E(2.62)}"/><a:ext cx="{E(7.52)}" cy="{E(1.50)}"/></a:xfrm>'))
        tx = sp.find(qn("p:txBody"))
        for p in tx.findall(qn("a:p")):
            tx.remove(p)
        bp = tx.find(qn("a:bodyPr"))
        for k, v in dict(lIns="0", tIns="0", rIns="0", bIns="0", anchor="t").items():
            bp.set(k, v)
        tx.append(etree.fromstring(
            f'<a:p xmlns:a="{NS_A}"><a:pPr algn="r" rtl="1"><a:lnSpc><a:spcPts val="5200"/></a:lnSpc></a:pPr>'
            f'<a:r><a:rPr lang="ar-SA" sz="4000" b="1" dirty="0" noProof="1"><a:solidFill><a:srgbClr val="{WHITE}"/></a:solidFill>'
            f'<a:latin typeface="{F_SEMI}" pitchFamily="2" charset="-78"/><a:cs typeface="{F_SEMI}" pitchFamily="2" charset="-78"/></a:rPr>'
            f'<a:t>{escape(title)}</a:t></a:r></a:p>'))
        sl.text("اسم البرنامج", 4.9, 2.08, 7.52, 0.44,
                [para(program, sz=18, color=TEAL, font=F_SEMI, pitch=26)], anchor="b")
        sl.text("العنوان الفرعي", 4.9, 4.22, 7.52, 0.62,
                [para(subtitle, sz=26, color=WHITE, font=F_MED, pitch=36)], anchor="t")
        sl.text("المُعِد والتاريخ", 4.9, 4.98, 7.52, 0.40,
                [para(byline, sz=16, color="D5D3E8", font=F_MED, pitch=24)], anchor="t")
        sl.text("التصنيف", 5.17, 6.95, 3.0, 0.36,
                [para(CLASSIFICATION, sz=12, color=GREEN, font=F_SEMI, algn="ctr", pitch=16)], anchor="ctr")
        self.slides.append(dict(sl=sl, kind="cover", title=title, label=label))
        return sl

    def divider(self, num, title, sub, label=""):
        slide = self.prs.slides.add_slide(self.L_DIVIDER)
        sl = Sl(slide, self.icons, label=label or f"divider {num}")
        sp = self._fill_ph(slide, ph_type="title", text=title)
        spPr = sp.find(qn("p:spPr"))
        for ch in list(spPr):
            spPr.remove(ch)
        spPr.append(etree.fromstring(
            f'<a:xfrm xmlns:a="{NS_A}"><a:off x="{E(4.9)}" y="{E(3.55)}"/><a:ext cx="{E(7.52)}" cy="{E(1.25)}"/></a:xfrm>'))
        tx = sp.find(qn("p:txBody"))
        for p in tx.findall(qn("a:p")):
            tx.remove(p)
        bp = tx.find(qn("a:bodyPr"))
        for k, v in dict(lIns="0", tIns="0", rIns="0", bIns="0", anchor="t").items():
            bp.set(k, v)
        tx.append(etree.fromstring(
            f'<a:p xmlns:a="{NS_A}"><a:pPr algn="r" rtl="1"><a:lnSpc><a:spcPts val="4800"/></a:lnSpc></a:pPr>'
            f'<a:r><a:rPr lang="ar-SA" sz="3600" b="1" dirty="0" noProof="1"><a:solidFill><a:srgbClr val="{WHITE}"/></a:solidFill>'
            f'<a:latin typeface="{F_SEMI}" pitchFamily="2" charset="-78"/><a:cs typeface="{F_SEMI}" pitchFamily="2" charset="-78"/></a:rPr>'
            f'<a:t>{escape(title)}</a:t></a:r></a:p>'))
        d = 1.15
        if isinstance(num, str) and num.startswith("Pi"):
            sl.oval("رمز القسم", R_EDGE - d, 2.28, d, fill=GREEN)
            sl.icon(num, WHITE, R_EDGE - d + 0.3, 2.28 + 0.3, d - 0.6)
        else:
            sl.box("رقم القسم", R_EDGE - d, 2.28, d, d,
                   [para(str(num), sz=46, color=WHITE, font=F_SEMI, b=True, algn="ctr", pitch=60)],
                   geom="ellipse", fill=GREEN, anchor="ctr", ins=(0, 0, 0, 0), fit=False)
        sl.text("وصف القسم", 4.9, 4.92, 7.52, 0.9,
                [para(sub, sz=20, color="D5D3E8", font=F_MED, pitch=30)], anchor="t")
        sl.text("التصنيف", 5.17, 6.95, 3.0, 0.36,
                [para(CLASSIFICATION, sz=12, color=GREEN, font=F_SEMI, algn="ctr", pitch=16)], anchor="ctr")
        self.slides.append(dict(sl=sl, kind="divider", title=title, label=sl.label))
        return sl

    # ------------------------------------------------------------------ notes / flags
    def notes(self, sl, paragraphs):
        ns = sl.slide.notes_slide
        tf = ns.notes_text_frame
        if tf is None:
            return
        tx = tf._txBody
        for p in tx.findall(qn("a:p")):
            tx.remove(p)
        for t in paragraphs:
            if t is None:
                continue
            bold = t.startswith("##")
            t = t.lstrip("#").strip() if bold else t
            tx.append(etree.fromstring(
                f'<a:p xmlns:a="{NS_A}"><a:pPr algn="r" rtl="1"/><a:r><a:rPr lang="ar-SA" sz="1200" b="{int(bold)}" dirty="0"/>'
                f'<a:t>{escape(t)}</a:t></a:r></a:p>'))

    def hide(self, sl):
        sl.slide._element.set("show", "0")

    def custom_show(self, name, labels):
        """named subset of slides (PowerPoint: Slide Show > Custom Slide Show)"""
        pres = self.prs.part._element
        lst = pres.find(qn("p:sldIdLst"))
        pos = {x["label"]: i for i, x in enumerate(self.slides)}
        rids = [lst[pos[l]].get("{%s}id" % NS_R) for l in labels]
        xml = (f'<p:custShowLst {NSDECL}><p:custShow name="{escape(name)}" id="0"><p:sldLst>'
               + "".join(f'<p:sld r:id="{r}"/>' for r in rids) + "</p:sldLst></p:custShow></p:custShowLst>")
        anchor = pres.find(qn("p:embeddedFontLst"))
        if anchor is None:
            anchor = pres.find(qn("p:notesSz"))
        anchor.addnext(etree.fromstring(xml))

    # ------------------------------------------------------------------ finalize
    def finalize(self, path, title, subject, author, keywords="", thumb_jpeg=None, hidden=0):
        cp = self.prs.core_properties
        cp.title, cp.subject, cp.author, cp.last_modified_by = title, subject, author, author
        cp.keywords, cp.comments = keywords, ""
        cp.category = ""
        now = datetime.datetime(2026, 10, 4, 12, 0, 0)
        cp.created, cp.modified, cp.revision = now, now, 1
        n = len(self.prs.slides)
        for part in self.prs.part.package.iter_parts():
            pn = str(part.partname)
            if pn == "/docProps/custom.xml":      # drop the template's inherited MSIP sensitivity label
                part._blob = (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                              b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
                              b'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"/>')
            elif pn == "/docProps/app.xml":
                part._blob = (
                    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
                    'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
                    f'<TotalTime>0</TotalTime><Application>Microsoft Office PowerPoint</Application>'
                    f'<PresentationFormat>Widescreen</PresentationFormat><Slides>{n}</Slides><Notes>{n}</Notes>'
                    f'<HiddenSlides>{hidden}</HiddenSlides><Company>هيئة الحكومة الرقمية</Company></Properties>').encode("utf-8")
        self.prs.save(path)
        if thumb_jpeg and os.path.exists(thumb_jpeg):
            self._swap_thumbnail(path, thumb_jpeg)

    @staticmethod
    def _swap_thumbnail(path, thumb):
        tmp = path + ".tmp"
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "docProps/thumbnail.jpeg":
                    data = open(thumb, "rb").read()
                zout.writestr(item, data)
        os.replace(tmp, path)
